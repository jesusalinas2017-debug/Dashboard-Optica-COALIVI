# app.py
# Dashboard ejecutivo · Venta particular · Óptica COALIVI
# Elaborado por MDF Consulting para la Dirección Ejecutiva de COALIVI (2026)
#
# Foco: solo venta particular (OT con convenio PARTICULAR). Licitaciones y convenios quedan fuera.
# Para correrlo en el computador:  streamlit run app.py

import base64
import glob
import io
import os
import zipfile

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import plotly.io as pio
import streamlit as st

import procesamiento as pr


# =============================================================================
# 1. Configuración general y colores COALIVI
# =============================================================================
CARPETA_APP = os.path.dirname(os.path.abspath(__file__))
RUTA_LOGO = os.path.join(CARPETA_APP, 'logo.png')

st.set_page_config(page_title='Dashboard Ejecutivo · Óptica COALIVI',
                   page_icon=RUTA_LOGO if os.path.exists(RUTA_LOGO) else '👓',
                   layout='wide', initial_sidebar_state='expanded')

AZUL = '#1A4474'        # azul institucional COALIVI
NARANJO = '#FA9219'     # naranjo institucional COALIVI
AZUL_SUAVE = '#A9B8CE'  # para el "año anterior" en las comparaciones
TINTA = '#1F2A44'
GRIS = '#6B7385'
GRILLA = '#E6E9EF'

# Paleta para categorías (validada para daltonismo, en este orden)
SERIES = ['#2F6DB5', '#F08C1A', '#1BAF7A', '#7B6FD6', '#E87BA4', '#EDA100']

# Colores fijos por entidad: cada línea / familia / tipo de derivador tiene siempre el mismo color
COLOR_LINEA = dict(zip(pr.LINEAS, SERIES))
COLOR_FAMILIA = {'Cristal': SERIES[0], 'Armazón': SERIES[1], 'Gafas': SERIES[2], 'Accesorio': SERIES[3]}
COLOR_DERIVADOR = {'Oftalmólogo/a': SERIES[0], 'Tecnólogo/a médico': SERIES[1], 'Otro': SERIES[3],
                   'Sin registro': '#C3C8D4'}
ORDEN_DERIVADOR = ['Oftalmólogo/a', 'Tecnólogo/a médico', 'Otro', 'Sin registro']

# Semáforo: solo para metas, siempre con texto (nunca solo color)
SEMAFORO = {'Cumple': ('#0CA30C', 'white'), 'En riesgo': ('#FAB219', TINTA),
            'No cumple': ('#D03B3B', 'white'), 'Sin dato': ('#C3C8D4', TINTA)}
ROJO = '#D03B3B'

MESES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
DIAS = ['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado', 'Domingo']

# Plantilla de Plotly con la identidad de COALIVI
pio.templates['coalivi'] = go.layout.Template(layout=dict(
    font=dict(family='Inter, Segoe UI, system-ui, sans-serif', size=13, color=TINTA),
    paper_bgcolor='white', plot_bgcolor='white',
    colorway=SERIES,
    separators=',.',                       # 1.234.567,8 (formato chileno)
    margin=dict(l=10, r=10, t=40, b=10),
    title=dict(font=dict(size=15, color=AZUL), x=0, xanchor='left'),
    xaxis=dict(showgrid=False, linecolor='#C3C8D4', ticks='', title=dict(font=dict(color=GRIS))),
    yaxis=dict(gridcolor=GRILLA, zeroline=False, linecolor='#C3C8D4', title=dict(font=dict(color=GRIS))),
    legend=dict(orientation='h', yanchor='bottom', y=1.0, xanchor='right', x=1, title=dict(text='')),
    hoverlabel=dict(bgcolor='white', bordercolor=AZUL, font=dict(color=TINTA)),
    bargap=0.25,
))
pio.templates.default = 'coalivi'


# =============================================================================
# 2. Estilo visual (CSS)
# =============================================================================
st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
html, body, [class*="css"] {{ font-family: 'Inter', 'Segoe UI', sans-serif; }}
.block-container {{ padding-top: 1.2rem; padding-bottom: 2rem; max-width: 1400px; }}

/* Encabezado */
.encabezado {{ background: linear-gradient(90deg, {AZUL} 0%, #24578F 100%);
  border-bottom: 5px solid {NARANJO}; border-radius: 14px; padding: 18px 26px; margin-bottom: 16px;
  display: flex; align-items: center; gap: 22px; }}
.encabezado img {{ height: 84px; width: 84px; }}
.encabezado h1 {{ color: white; font-size: 1.75rem; font-weight: 700; margin: 0; padding: 0; }}
.encabezado p {{ color: #D9E2EF; margin: 4px 0 0 0; font-size: 0.92rem; }}
.encabezado .marca {{ color: {NARANJO}; font-weight: 700; letter-spacing: 1px; font-size: 0.78rem; }}

/* Tarjetas KPI */
.kpi {{ background: white; border-radius: 12px; padding: 14px 16px 12px 16px;
  border-left: 5px solid {NARANJO}; box-shadow: 0 1px 3px rgba(26,68,116,0.12); min-height: 128px; }}
.kpi-t {{ color: {GRIS}; font-size: 0.76rem; font-weight: 600; text-transform: uppercase; letter-spacing: 0.4px; }}
.kpi-v {{ color: {AZUL}; font-size: 1.6rem; font-weight: 700; line-height: 1.25; margin-top: 2px; }}
.kpi-d {{ font-size: 0.8rem; font-weight: 600; margin-top: 2px; }}
.kpi-d span {{ color: {GRIS}; font-weight: 400; }}
.kpi-m {{ font-size: 0.75rem; color: {GRIS}; margin-top: 6px; }}
.sube {{ color: #0A7A0A; }}  .baja {{ color: {ROJO}; }}  .neutro {{ color: {GRIS}; }}
.chip {{ display: inline-block; padding: 1px 8px; border-radius: 10px; font-size: 0.72rem; font-weight: 700;
  margin-right: 6px; white-space: nowrap; }}

/* Títulos de sección */
.seccion {{ color: {AZUL}; font-size: 1.1rem; font-weight: 700; border-left: 4px solid {NARANJO};
  padding-left: 10px; margin: 22px 0 6px 0; }}
.subtitulo {{ color: {GRIS}; font-size: 0.85rem; margin: -2px 0 8px 14px; }}

/* Caja de hallazgos y nota de método */
.hallazgos {{ background: #FFF6EA; border: 1px solid #F8D3A0; border-radius: 12px; padding: 14px 20px; }}
.hallazgos h4 {{ color: {AZUL}; margin: 0 0 6px 0; font-size: 1rem; }}
.hallazgos li {{ margin-bottom: 4px; color: {TINTA}; font-size: 0.92rem; }}
.nota {{ background: #EEF3FA; border-left: 4px solid {AZUL}; border-radius: 8px; padding: 10px 14px;
  font-size: 0.85rem; color: {TINTA}; margin: 6px 0 12px 0; }}

/* Tablas ejecutivas */
.tabla {{ width: 100%; border-collapse: collapse; font-size: 0.88rem; background: white; }}
.tabla th {{ background: {AZUL}; color: white; text-align: right; padding: 8px 10px; font-weight: 600; }}
.tabla th:first-child, .tabla td:first-child {{ text-align: left; }}
.tabla td {{ border-bottom: 1px solid {GRILLA}; padding: 7px 10px; color: {TINTA}; text-align: right; }}
.tabla tr.total td {{ font-weight: 700; border-top: 2px solid {AZUL}; background: #F4F6FA; }}
.tabla td.neg {{ color: {ROJO}; }}

/* Pestañas */
.stTabs [data-baseweb="tab-list"] {{ gap: 4px; border-bottom: 2px solid {GRILLA}; }}
.stTabs [data-baseweb="tab"] {{ padding: 8px 16px; font-weight: 600; color: {AZUL}; }}

/* Barra lateral */
[data-testid="stSidebar"] .logo {{ text-align: center; margin-bottom: 6px; }}
[data-testid="stSidebar"] .logo img {{ width: 110px; }}
[data-testid="stSidebar"] .logo-texto {{ background: {AZUL}; color: white; border-radius: 10px; padding: 14px;
  text-align: center; border-bottom: 4px solid {NARANJO}; font-weight: 700; letter-spacing: 2px; }}
.pie {{ color: {GRIS}; font-size: 0.8rem; text-align: center; margin-top: 30px; }}
</style>
""", unsafe_allow_html=True)


# =============================================================================
# 3. Funciones de formato y de dibujo
# =============================================================================
def miles(x, dec=0):
    # 1234567.8 -> '1.234.568' (separador de miles chileno)
    if x is None or pd.isna(x):
        return '–'
    texto = f'{x:,.{dec}f}'
    return texto.replace(',', 'X').replace('.', ',').replace('X', '.')


def etiqueta_mes(fechas):
    # 2025-03-01 -> 'Mar 25' (meses en español para los ejes)
    return fechas.dt.month.map(lambda m: MESES[int(m) - 1]) + ' ' + fechas.dt.strftime('%y')


def clp(x):
    return '–' if x is None or pd.isna(x) else ('-$' if x < 0 else '$') + miles(abs(x))


def mm(x):
    # Montos grandes en millones: $245,3 MM
    if x is None or pd.isna(x):
        return '–'
    signo = '-' if x < 0 else ''
    return signo + '$' + miles(abs(x) / 1e6, 1) + ' MM' if abs(x) >= 1e6 else clp(x)


def pct(x, dec=1):
    return '–' if x is None or pd.isna(x) else miles(x, dec) + '%'


def variacion(actual, anterior):
    # Variación porcentual; None si no hay base de comparación
    if anterior is None or pd.isna(anterior) or anterior == 0 or actual is None or pd.isna(actual):
        return None
    return 100 * (actual - anterior) / abs(anterior)


def diferencia(actual, anterior):
    # Diferencia en puntos porcentuales; None si falta alguno
    if actual is None or anterior is None or pd.isna(actual) or pd.isna(anterior):
        return None
    return actual - anterior


def estado_meta(valor, meta, mayor_es_mejor=True):
    # Semáforo: Cumple / En riesgo (a menos de 10% de la meta) / No cumple
    if meta is None or valor is None or pd.isna(valor):
        return 'Sin dato'
    holgura = max(abs(meta) * 0.10, 1)
    brecha = (valor - meta) if mayor_es_mejor else (meta - valor)
    if brecha >= 0:
        return 'Cumple'
    return 'En riesgo' if brecha >= -holgura else 'No cumple'


def chip(estado, texto=None):
    fondo, letra = SEMAFORO[estado]
    return f'<span class="chip" style="background:{fondo};color:{letra}">● {texto or estado}</span>'


def tarjeta(titulo, valor, delta=None, comparado='', invertir=False, puntos=False, meta=None):
    # Tarjeta KPI en HTML.
    # invertir = True cuando subir es malo (ej. % de errores); puntos = delta en puntos porcentuales
    # meta = (estado, texto) para mostrar el semáforo
    if delta is None:
        linea = f'<div class="kpi-d neutro">{comparado}</div>' if comparado else '<div class="kpi-d">&nbsp;</div>'
    else:
        bueno = (delta >= 0) != invertir
        clase = 'neutro' if abs(delta) < 0.05 else ('sube' if bueno else 'baja')
        flecha = '▲' if delta > 0 else ('▼' if delta < 0 else '●')
        unidad = ' pp' if puntos else '%'
        linea = (f'<div class="kpi-d {clase}">{flecha} {miles(abs(delta), 1)}{unidad} '
                 f'<span>{comparado}</span></div>')
    pie = f'<div class="kpi-m">{chip(meta[0])}{meta[1]}</div>' if meta else ''
    return (f'<div class="kpi"><div class="kpi-t">{titulo}</div><div class="kpi-v">{valor}</div>'
            f'{linea}{pie}</div>')


def fila_kpi(tarjetas):
    columnas = st.columns(len(tarjetas), gap='small')
    for col, html in zip(columnas, tarjetas):
        col.markdown(html, unsafe_allow_html=True)


def seccion(texto, sub=''):
    st.markdown(f'<div class="seccion">{texto}</div>', unsafe_allow_html=True)
    if sub:
        st.markdown(f'<div class="subtitulo">{sub}</div>', unsafe_allow_html=True)


def nota(texto):
    st.markdown(f'<div class="nota">{texto}</div>', unsafe_allow_html=True)


def mostrar(fig, alto=360):
    fig.update_layout(height=alto)
    st.plotly_chart(fig, config={'displaylogo': False})


def sin_datos(texto='No hay datos para el período seleccionado.'):
    st.info(texto)


def barras_h(df, x, y, titulo, color=AZUL, formato='$', alto=360, etiqueta=None):
    # Barras horizontales de una sola serie (un solo color), ordenadas de mayor a menor
    df = df.sort_values(x)
    fig = px.bar(df, x=x, y=y, orientation='h', title=titulo)
    fig.update_traces(marker_color=color, marker_line_width=0)
    if etiqueta is not None:
        fig.update_traces(text=df[etiqueta], textposition='outside', cliponaxis=False)
    plantilla = '%{y}<br>' + ('$%{x:,.0f}' if formato == '$' else '%{x:,.1f}' + formato) + '<extra></extra>'
    fig.update_traces(hovertemplate=plantilla)
    fig.update_layout(xaxis_title=None, yaxis_title=None,
                      xaxis=dict(showticklabels=etiqueta is None, gridcolor=GRILLA, showgrid=True))
    if etiqueta is not None and len(df) and df[x].max() > 0:
        fig.update_xaxes(range=[0, df[x].max() * 1.22])   # espacio para que no se corte la etiqueta
    mostrar(fig, alto)


def actual_vs_anterior(df, x, y, nombre_act, nombre_ant, titulo, formato='$', alto=360, orden_x=None):
    # Barras agrupadas: período actual (azul COALIVI) vs mismo período del año anterior
    fig = px.bar(df, x=x, y=y, color='Período', barmode='group', title=titulo,
                 color_discrete_map={nombre_act: AZUL, nombre_ant: AZUL_SUAVE},
                 category_orders={'Período': [nombre_ant, nombre_act], **({x: orden_x} if orden_x else {})})
    if formato == '$':
        valor = '$%{y:,.0f}'
    elif formato == '%':
        valor = '%{y:,.1f}%'
    else:
        valor = '%{y:,.0f}' + formato
    fig.update_traces(hovertemplate='%{x}<br>' + valor + '<extra>%{fullData.name}</extra>',
                      marker_line_width=0)
    fig.update_layout(xaxis_title=None, yaxis_title=None)
    mostrar(fig, alto)


def linea_mensual(df, y, formato='$', alto=320, meta=None, color=AZUL):
    # Línea mes a mes con etiquetas de mes en español
    df = df.sort_values('mes').copy()
    df['Mes'] = etiqueta_mes(df['mes'])
    fig = px.line(df, x='Mes', y=y, markers=True)
    valor = '$%{y:,.0f}' if formato == '$' else '%{y:,.1f}' + formato
    fig.update_traces(line=dict(color=color, width=2), marker=dict(size=7, color=color),
                      hovertemplate='%{x}<br>' + valor + '<extra></extra>')
    if meta is not None:
        fig.add_hline(y=meta, line_color=GRIS, line_width=1, annotation_text='meta', annotation_position='top left')
    fig.update_layout(xaxis_title=None, yaxis_title=None, xaxis_type='category')
    if formato == '%':
        fig.update_yaxes(rangemode='tozero', ticksuffix='%')
    mostrar(fig, alto)


def tabla_html(df, con_total=False, negativas=()):
    # Tabla ejecutiva en HTML. Si con_total=True, la última fila se destaca como total
    cab = ''.join(f'<th>{c}</th>' for c in df.columns)
    filas = ''
    for i, (_, fila) in enumerate(df.iterrows()):
        clase = ' class="total"' if con_total and i == len(df) - 1 else ''
        celdas = ''
        for c in df.columns:
            valor = fila[c]
            neg = ' class="neg"' if c in negativas and str(valor).startswith('-') else ''
            celdas += f'<td{neg}>{valor}</td>'
        filas += f'<tr{clase}>{celdas}</tr>'
    st.markdown(f'<table class="tabla"><tr>{cab}</tr>{filas}</table>', unsafe_allow_html=True)


def logo_base64():
    if os.path.exists(RUTA_LOGO):
        with open(RUTA_LOGO, 'rb') as f:
            return base64.b64encode(f.read()).decode()
    return None


# =============================================================================
# 4. Carga de datos
# =============================================================================
@st.cache_data(show_spinner='Cargando datos de COALIVI...')
def cargar_repositorio():
    # Lee las tablas limpias (.csv.gz) que vienen en el repositorio.
    # Las busca en data/, en la carpeta principal o en cualquier subcarpeta,
    # por si al subir a GitHub la carpeta data/ quedó con otro nombre o se aplanó.
    tablas = {}
    for clave in pr.FECHAS:
        encontrados = sorted(glob.glob(os.path.join(CARPETA_APP, '**', f'{clave}.csv.gz'), recursive=True))
        if encontrados:
            tablas[clave] = pr.leer_tabla(encontrados[0], clave)

    # Si se subió el .zip tal cual a GitHub, se leen las tablas desde adentro del zip
    for ruta_zip in glob.glob(os.path.join(CARPETA_APP, '**', '*.zip'), recursive=True):
        with zipfile.ZipFile(ruta_zip) as z:
            for nombre in z.namelist():
                clave = os.path.basename(nombre).replace('.csv.gz', '')
                if nombre.endswith('.csv.gz') and clave in pr.FECHAS and clave not in tablas:
                    tablas[clave] = pr.leer_tabla(io.BytesIO(z.read(nombre)), clave)
    return tablas


@st.cache_data(show_spinner='Procesando planilla subida...')
def procesar_subida(nombre, contenido):
    # Limpia una planilla original subida desde la barra lateral
    return pr.procesar_archivo(contenido, nombre)


datos = cargar_repositorio()
origen_datos = {clave: 'Repositorio' for clave in datos}
LOGO = logo_base64()

# ---- Barra lateral: logo y carga de archivos
with st.sidebar:
    if LOGO:
        st.markdown(f'<div class="logo"><img src="data:image/png;base64,{LOGO}"></div>', unsafe_allow_html=True)
    else:
        st.markdown('<div class="logo-texto">COALIVI</div>', unsafe_allow_html=True)

    with st.expander('Cargar / actualizar bases de datos', expanded=not datos):
        st.caption('Sube las planillas originales (.xlsx) tal como las entrega TI: boletas, OT de sala de ventas '
                   'e informe de costos de sala. Reemplazan a las del repositorio mientras la página esté abierta.')
        subidos = st.file_uploader('Planillas Excel', type=['xlsx'],
                                   accept_multiple_files=True, label_visibility='collapsed')
        for archivo in subidos or []:
            try:
                clave, tabla = procesar_subida(archivo.name, archivo.getvalue())
                datos[clave] = tabla
                origen_datos[clave] = f'Subido: {archivo.name}'
                st.success(f'{pr.BASES[clave]["nombre"]}: {miles(len(tabla))} filas')
            except Exception as error:
                st.warning(f'{archivo.name}: {error}')

if 'ventas' not in datos or 'ot_sala' not in datos:
    st.markdown('<div class="encabezado"><div><span class="marca">ÓPTICA COALIVI · MDF CONSULTING</span>'
                '<h1>Dashboard Ejecutivo · Venta particular</h1>'
                '<p>Sube las planillas en la barra lateral para comenzar.</p></div></div>', unsafe_allow_html=True)
    st.warning('No se encontraron los datos en el repositorio. Revisa en GitHub que exista la carpeta **data/** '
               'con ventas.csv.gz, ot_sala.csv.gz y costos.csv.gz. Mientras tanto, puedes subir las planillas '
               'originales en la barra lateral.')
    st.stop()

# Tablas vacías para las bases que no estén cargadas (así la app no se cae)
for clave in pr.FECHAS:
    if clave not in datos:
        datos[clave] = pd.DataFrame(columns=['fecha', 'ot', 'anio', 'mes'])
        origen_datos[clave] = 'No cargada'

# ---- Filtro a venta particular: la OT define qué es particular
ventas_sala, ot_sala, costos = pr.solo_particular(datos['ventas'], datos['ot_sala'], datos['costos'])
ventas_todas = ventas_sala[ventas_sala['es_particular']]       # incluye las boletas anuladas (NC)
ventas = ventas_todas[~ventas_todas['es_nc']]                   # venta válida

# ---- Filtros de período, gastos y metas
fecha_max = ventas['fecha'].max()
anios = sorted(ventas['anio'].dropna().astype(int).unique().tolist(), reverse=True)

with st.sidebar:
    st.markdown('**Período de análisis**')
    anio = st.selectbox('Año', anios, index=0)
    mes_ini, mes_fin = st.select_slider('Meses', options=MESES, value=(MESES[0], MESES[-1]))
    st.caption('Todo se compara con el **mismo tramo del año anterior**.')

    with st.expander('Gastos operacionales (pendiente)'):
        st.caption('Gasto **mensual** atribuible a la venta particular, sin IVA. Todavía no tenemos las cifras de '
                   'COALIVI: mientras estén en 0, el dashboard llega hasta el margen directo.')
        gasto_rem = st.number_input('Remuneraciones equipo óptica', min_value=0, value=0, step=100_000)
        gasto_arr = st.number_input('Arriendo y gastos comunes', min_value=0, value=0, step=100_000)
        gasto_ser = st.number_input('Servicios, sistemas y comunicaciones', min_value=0, value=0, step=50_000)
        gasto_mkt = st.number_input('Marketing y difusión', min_value=0, value=0, step=50_000)
        gasto_otr = st.number_input('Otros gastos de administración', min_value=0, value=0, step=50_000)
        con_iva = st.checkbox('Los precios de venta incluyen IVA (19%)', value=True)

    with st.expander('Metas'):
        meta_crec = st.number_input('Crecimiento de ingresos vs año anterior (%)', value=5.0, step=1.0)
        meta_md = st.number_input('Margen directo mínimo (%)', value=70.0, step=1.0)
        meta_mo = st.number_input('Margen operacional mínimo (%)', value=20.0, step=1.0)
        meta_der = st.number_input('OT con derivador registrado mínimo (%)', value=80.0, step=5.0)
        meta_lt = st.number_input('Plazo de entrega máximo (días, mediana)', value=7.0, step=1.0)
        meta_atr = st.number_input('Entregas fuera de plazo máximo (%)', value=10.0, step=1.0)
        meta_err = st.number_input('Venta con OT mal registrada máximo (%)', value=2.0, step=0.5)

GASTO_MENSUAL = gasto_rem + gasto_arr + gasto_ser + gasto_mkt + gasto_otr
HAY_GASTOS = GASTO_MENSUAL > 0
FACTOR_IVA = 1.19 if con_iva else 1.0

i1, i2 = MESES.index(mes_ini) + 1, MESES.index(mes_fin) + 1
ini = pd.Timestamp(int(anio), i1, 1)
fin = pd.Timestamp(int(anio), i2, 1) + pd.offsets.MonthEnd(0)
if ini <= fecha_max < fin:
    fin = fecha_max                        # el año en curso se corta en el último dato
ini_ant = ini - pd.DateOffset(years=1)
fin_ant = fin - pd.DateOffset(years=1)
MESES_PERIODO = ((fin - ini).days + 1) / 30.44   # meses del período (para los gastos)
MESES_RANGO = range(i1, fin.month + 1)            # meses con datos (para los gráficos mensuales)

ACT, ANT = str(int(anio)), str(int(anio) - 1)
COMPARADO = f'vs mismo tramo {ANT}'


def periodo(df, anterior=False, col='fecha'):
    # Filtra una tabla al período elegido (o al mismo tramo del año anterior)
    if df.empty or col not in df.columns:
        return df
    a, b = (ini_ant, fin_ant) if anterior else (ini, fin)
    return df[(df[col] >= a) & (df[col] <= b)]


with st.sidebar:
    st.caption(f'Período: {ini:%d-%m-%Y} al {fin:%d-%m-%Y}')
    st.markdown('<div class="pie">Elaborado por <b>MDF Consulting</b><br>'
                'Ingeniería Civil Industrial · UDD · 2026</div>', unsafe_allow_html=True)

# Tablas del período actual y del anterior
v_act, v_ant = periodo(ventas), periodo(ventas, True)
if v_act.empty:
    st.warning(f'No hay ventas particulares entre {ini:%d-%m-%Y} y {fin:%d-%m-%Y}. '
               'Cambia el año o los meses en la barra lateral.')
    st.stop()
s_act, s_ant = periodo(ot_sala), periodo(ot_sala, True)
vs_act, vs_ant = periodo(ventas_sala), periodo(ventas_sala, True)      # toda la sala (calidad de datos)

# Si no hay ventas el año anterior (2024 es el primer año), no se compara nada
if v_ant.empty:
    s_ant, vs_ant = s_ant.iloc[0:0], vs_ant.iloc[0:0]
    COMPARADO = f'sin datos de {ANT} para comparar'


# =============================================================================
# 5. Estado de resultados por línea de producto
# =============================================================================
def porcentaje_costo():
    # % de costo directo sobre el precio de venta (con IVA), por línea, según el informe de costos.
    # Las líneas sin datos propios usan el promedio general.
    ok = costos[costos['costeable']] if not costos.empty else costos
    if not ok.empty and ok['venta'].sum() > 0:
        general = ok['costo'].sum() / ok['venta'].sum()
        por_linea = ok.groupby('linea').agg(venta=('venta', 'sum'), costo=('costo', 'sum'), n=('venta', 'size'))
    else:
        general, por_linea = 0.25, pd.DataFrame(columns=['venta', 'costo', 'n'])
    filas = []
    for linea in pr.LINEAS:
        if linea in por_linea.index and por_linea.loc[linea, 'n'] >= 20:
            r = por_linea.loc[linea]
            filas.append({'linea': linea, 'pct_costo': r['costo'] / r['venta'],
                          'fuente': f'Informe de costos ({miles(r["n"])} productos)'})
        else:
            filas.append({'linea': linea, 'pct_costo': general, 'fuente': 'Promedio general (sin datos propios)'})
    return pd.DataFrame(filas).set_index('linea')


PCT_COSTO = porcentaje_costo()


def estado_resultados(v, meses):
    # Ingresos netos, costo directo, margen directo, gastos operacionales y margen operacional por línea
    bruto = v.groupby('linea')['total'].sum().reindex(pr.LINEAS, fill_value=0)
    er = pd.DataFrame(index=pr.LINEAS)
    er['ingresos'] = bruto / FACTOR_IVA
    er['costo'] = bruto * PCT_COSTO['pct_costo']          # el % de costo es sobre el precio de venta
    er['margen_directo'] = er['ingresos'] - er['costo']
    total_ing = er['ingresos'].sum()
    participacion = er['ingresos'] / total_ing if total_ing else 0
    er['gastos'] = GASTO_MENSUAL * meses * participacion   # gastos repartidos según los ingresos
    er['margen_operacional'] = er['margen_directo'] - er['gastos']
    er.loc['Total'] = er.sum()
    er['pct_ingresos'] = 100 * er['ingresos'] / total_ing if total_ing else np.nan
    er['pct_md'] = 100 * er['margen_directo'] / er['ingresos'].replace(0, np.nan)
    er['pct_mo'] = 100 * er['margen_operacional'] / er['ingresos'].replace(0, np.nan)
    if not HAY_GASTOS:                                     # sin gastos no hay margen operacional
        er[['gastos', 'margen_operacional', 'pct_mo']] = np.nan
    return er


ER_ACT = estado_resultados(v_act, MESES_PERIODO)
ER_ANT = estado_resultados(v_ant, MESES_PERIODO)
TOT, TOT_ANT = ER_ACT.loc['Total'], ER_ANT.loc['Total']


# =============================================================================
# 6. Indicadores comunes y metas
# =============================================================================
def resumen_ventas(v):
    venta = v['total'].sum()
    boletas = v['folio'].nunique()
    return venta, boletas, (venta / boletas if boletas else np.nan)


def unidades(v):
    # Una venta a crédito se boletea en dos líneas (abono + saldo) con la misma cantidad,
    # así que las unidades se cuentan sin las líneas de saldo
    return v.loc[v['clase_linea'] != 'Saldo', 'cantidad'].sum()


def pct_error(v):
    # % del monto facturado cuya OT está mal o no registrada
    total = v['total'].sum() if not v.empty else 0
    return 100 * v.loc[v['estado_ot'] != 'Registrada', 'total'].sum() / total if total else np.nan


def ticket_ot(s):
    con_valor = s[s['total'] > 0]['total'] if not s.empty else pd.Series(dtype=float)
    return con_valor.mean() if len(con_valor) else np.nan


def pct_atraso(s):
    n = s['atraso'].notna().sum() if not s.empty else 0
    return 100 * (s['atraso'] > 0).sum() / n if n else np.nan


def pct_derivacion(s):
    return 100 * (s['origen_tipo'] != 'Sin registro').mean() if not s.empty else np.nan


def pct_descuento(s):
    return 100 * s['con_descuento'].mean() if not s.empty else np.nan


venta_act, boletas_act, ticket_act = resumen_ventas(v_act)
venta_ant, boletas_ant, ticket_ant = resumen_ventas(v_ant)
lt_act = s_act['lead_time'].median() if not s_act.empty else np.nan
lt_ant = s_ant['lead_time'].median() if not s_ant.empty else np.nan
err_act, err_ant = pct_error(vs_act), pct_error(vs_ant)
atr_act, atr_ant = pct_atraso(s_act), pct_atraso(s_ant)
der_act, der_ant = pct_derivacion(s_act), pct_derivacion(s_ant)
dsc_act, dsc_ant = pct_descuento(s_act), pct_descuento(s_ant)
crec_ing = variacion(TOT['ingresos'], TOT_ANT['ingresos'])

# Cuadro de metas: (indicador, valor actual, valor anterior, unidad, meta, mayor_es_mejor)
METAS = [
    ('Crecimiento de ingresos', crec_ing, None, '%', meta_crec, True),
    ('Margen directo', TOT['pct_md'], TOT_ANT['pct_md'], '%', meta_md, True),
    ('Margen operacional', TOT['pct_mo'], TOT_ANT['pct_mo'], '%', meta_mo, True),
    ('OT con derivador registrado', der_act, der_ant, '%', meta_der, True),
    ('Plazo de entrega (mediana)', lt_act, lt_ant, 'días', meta_lt, False),
    ('Entregas fuera de plazo', atr_act, atr_ant, '%', meta_atr, False),
    ('Venta con OT mal registrada', err_act, err_ant, '%', meta_err, False),
]


def texto_meta(meta, unidad, mayor_es_mejor):
    signo = '≥' if mayor_es_mejor else '≤'
    return f'meta {signo} ' + (pct(meta) if unidad == '%' else f'{miles(meta)} días')


def semaforo(nombre):
    # Devuelve (estado, texto de la meta) para mostrar en la tarjeta
    for n, valor, _, unidad, meta, mayor in METAS:
        if n == nombre:
            return estado_meta(valor, meta, mayor), texto_meta(meta, unidad, mayor)
    return None


def valor_meta(nombre, valor, unidad):
    if nombre == 'Margen operacional' and not HAY_GASTOS:
        return 'pendiente'
    return pct(valor) if unidad == '%' else ('–' if valor is None or pd.isna(valor) else f'{miles(valor)} días')


# =============================================================================
# 7. Reporte ejecutivo descargable (Excel)
# =============================================================================
def reporte_excel():
    from openpyxl.styles import Alignment, Font, PatternFill

    salida = io.BytesIO()
    with pd.ExcelWriter(salida, engine='openpyxl') as w:
        # Hoja 1: indicadores principales y cuadro de metas
        principales = pd.DataFrame([
            ['Ingresos netos ($)', TOT['ingresos'], TOT_ANT['ingresos'], '', ''],
            ['Margen directo ($)', TOT['margen_directo'], TOT_ANT['margen_directo'], '', ''],
            ['Margen operacional ($)', TOT['margen_operacional'], TOT_ANT['margen_operacional'], '', ''],
            ['OT particulares', len(s_act), len(s_ant), '', ''],
            ['Ticket promedio por OT ($)', ticket_ot(s_act), ticket_ot(s_ant), '', ''],
            ['OT con descuento (%)', dsc_act, dsc_ant, '', ''],
        ] + [[f'{n} ({u})', valor, ant, meta, estado_meta(valor, meta, mayor)]
             for n, valor, ant, u, meta, mayor in METAS],
            columns=['Indicador', f'Actual ({ACT})', f'Mismo tramo {ANT}', 'Meta', 'Estado'])
        principales = principales.round(1)
        principales.to_excel(w, sheet_name='Resumen', startrow=4, index=False)

        # Hoja 2: estado de resultados por línea
        er = ER_ACT[['ingresos', 'pct_ingresos', 'costo', 'margen_directo', 'pct_md', 'gastos',
                     'margen_operacional', 'pct_mo']].reset_index()
        er.columns = ['Línea', 'Ingresos netos', '% de ingresos', 'Costo directo', 'Margen directo',
                      '% margen directo', 'Gastos operacionales', 'Margen operacional', '% margen operacional']
        er.round(1).to_excel(w, sheet_name='Resultados por línea', startrow=4, index=False)

        # Hoja 3: venta mensual
        vm = pd.DataFrame({'Mes': [MESES[m - 1] for m in MESES_RANGO],
                           f'Venta {ACT}': [v_act.loc[v_act['fecha'].dt.month == m, 'total'].sum() for m in MESES_RANGO],
                           f'Venta {ANT}': [v_ant.loc[v_ant['fecha'].dt.month == m, 'total'].sum() for m in MESES_RANGO]})
        vm.to_excel(w, sheet_name='Venta mensual', startrow=4, index=False)

        # Hoja 4: supuestos
        sup = pd.DataFrame({'Línea': PCT_COSTO.index, '% costo sobre precio': 100 * PCT_COSTO['pct_costo'].values,
                            'Fuente': PCT_COSTO['fuente'].values})
        sup.round(1).to_excel(w, sheet_name='Supuestos', startrow=4, index=False)
        gastos = pd.DataFrame({'Gasto operacional mensual': ['Remuneraciones', 'Arriendo y gastos comunes',
                                                             'Servicios y sistemas', 'Marketing', 'Otros', 'TOTAL'],
                               'Monto ($)': [gasto_rem, gasto_arr, gasto_ser, gasto_mkt, gasto_otr, GASTO_MENSUAL]})
        gastos.to_excel(w, sheet_name='Supuestos', startrow=len(sup) + 7, index=False)

        # Formato: título, encabezados azules, números y anchos de columna
        titulos = {'Resumen': 'Reporte ejecutivo · Venta particular · Óptica COALIVI',
                   'Resultados por línea': 'Estado de resultados por línea de producto (montos netos de IVA)',
                   'Venta mensual': 'Venta particular mensual (con IVA)',
                   'Supuestos': 'Supuestos del cálculo de márgenes'}
        for hoja, titulo in titulos.items():
            ws = w.book[hoja]
            ws['A1'] = titulo
            ws['A1'].font = Font(bold=True, size=14, color='1A4474')
            ws['A2'] = f'Período {ini:%d-%m-%Y} al {fin:%d-%m-%Y} · comparado con el mismo tramo de {ANT}'
            ws['A3'] = f'Generado el {pd.Timestamp.today():%d-%m-%Y} · MDF Consulting'
            ws['A3'].font = Font(italic=True, color='6B7385')
            for fila in ws.iter_rows(min_row=6):
                for celda in fila:
                    if isinstance(celda.value, (int, float)):
                        celda.number_format = '#,##0' if abs(celda.value) >= 1000 else '#,##0.0'
            for fila_cab in (5, len(sup) + 8) if hoja == 'Supuestos' else (5,):
                for celda in ws[fila_cab]:
                    if celda.value is not None:
                        celda.fill = PatternFill('solid', fgColor='1A4474')
                        celda.font = Font(bold=True, color='FFFFFF')
                        celda.alignment = Alignment(horizontal='center', wrap_text=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 20
            ws.column_dimensions['A'].width = 40
    return salida.getvalue()


# =============================================================================
# 8. Encabezado y pestañas
# =============================================================================
img_logo = f'<img src="data:image/png;base64,{LOGO}">' if LOGO else ''
st.markdown(f"""
<div class="encabezado">
  {img_logo}
  <div>
    <span class="marca">ÓPTICA COALIVI · CUADRO DE MANDO DE LA DIRECCIÓN EJECUTIVA</span>
    <h1>Dashboard Ejecutivo · Venta particular</h1>
    <p>Período {ini:%d-%m-%Y} al {fin:%d-%m-%Y} · comparado con el mismo tramo de {ANT} ·
       datos de ventas hasta el {fecha_max:%d-%m-%Y} · solo OT con convenio PARTICULAR</p>
  </div>
</div>""", unsafe_allow_html=True)

tabs = st.tabs(['Resumen ejecutivo', 'Resultados por línea', 'Comercial', 'Operaciones', 'Calidad de datos'])


# =============================================================================
# 9. Pestaña: Resumen ejecutivo
# =============================================================================
with tabs[0]:
    col_a, col_b = st.columns([4, 1])
    col_a.caption('Ingresos netos de IVA de la venta particular. Costo directo estimado con el informe de costos. '
                  + ('Gastos operacionales según la barra lateral.' if HAY_GASTOS
                     else 'Gastos operacionales: pendientes (el margen operacional se calcula cuando se ingresen).'))
    col_b.download_button('Descargar reporte (Excel)', reporte_excel(),
                          file_name=f'Reporte_COALIVI_particular_{ini:%Y%m%d}_{fin:%Y%m%d}.xlsx',
                          mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                          key='reporte_resumen')

    fila_kpi([
        tarjeta('Ingresos netos', mm(TOT['ingresos']), crec_ing, COMPARADO,
                meta=semaforo('Crecimiento de ingresos')),
        tarjeta('Margen directo', pct(TOT['pct_md']), None, f'{mm(TOT["margen_directo"])} · estimado',
                meta=semaforo('Margen directo')),
        tarjeta('Margen operacional', pct(TOT['pct_mo']) if HAY_GASTOS else 'Pendiente', None,
                f'{mm(TOT["margen_operacional"])} · estimado' if HAY_GASTOS else 'faltan los gastos de COALIVI',
                meta=semaforo('Margen operacional') if HAY_GASTOS else None),
        tarjeta('Ticket promedio por OT', clp(ticket_ot(s_act)),
                variacion(ticket_ot(s_act), ticket_ot(s_ant)), COMPARADO),
    ])
    st.write('')
    fila_kpi([
        tarjeta('OT particulares', miles(len(s_act)), variacion(len(s_act), len(s_ant)), COMPARADO),
        tarjeta('OT con derivador registrado', pct(der_act), diferencia(der_act, der_ant), COMPARADO,
                puntos=True, meta=semaforo('OT con derivador registrado')),
        tarjeta('Plazo de entrega (mediana)', f'{miles(lt_act)} días', variacion(lt_act, lt_ant), COMPARADO,
                invertir=True, meta=semaforo('Plazo de entrega (mediana)')),
        tarjeta('OT con descuento', pct(dsc_act), diferencia(dsc_act, dsc_ant), COMPARADO,
                invertir=True, puntos=True),
    ])

    col1, col2 = st.columns([3, 2], gap='medium')
    with col1:
        if HAY_GASTOS:
            seccion('Del ingreso al margen operacional', f'período {ACT} · montos netos de IVA')
            pasos = ['Ingresos netos', 'Costo directo', 'Margen directo', 'Gastos operacionales', 'Margen operacional']
            valores = [TOT['ingresos'], -TOT['costo'], TOT['margen_directo'], -TOT['gastos'], TOT['margen_operacional']]
            y = [valores[0], valores[1], 0, valores[3], 0]
            medidas = ['absolute', 'relative', 'total', 'relative', 'total']
        else:
            seccion('Del ingreso al margen directo', f'período {ACT} · montos netos de IVA · gastos operacionales pendientes')
            pasos = ['Ingresos netos', 'Costo directo', 'Margen directo']
            valores = [TOT['ingresos'], -TOT['costo'], TOT['margen_directo']]
            y = [valores[0], valores[1], 0]
            medidas = ['absolute', 'relative', 'total']
        fig = go.Figure(go.Waterfall(
            x=pasos, y=y, measure=medidas,
            text=[mm(x) for x in valores], textposition='outside',
            increasing=dict(marker=dict(color=AZUL)), decreasing=dict(marker=dict(color=NARANJO)),
            totals=dict(marker=dict(color=AZUL)), connector=dict(line=dict(color='#C3C8D4', width=1)),
            hovertemplate='%{x}<br>$%{y:,.0f}<extra></extra>'))
        tope = max(TOT['ingresos'], 1) * 1.15
        piso = min(0, TOT['margen_operacional'] * 1.3) if HAY_GASTOS else 0
        fig.update_layout(yaxis_title=None, showlegend=False, yaxis_range=[piso, tope])
        mostrar(fig, 360)
    with col2:
        seccion('Cuadro de metas', 'las metas se editan en la barra lateral')
        tabla_html(pd.DataFrame([{
            'Indicador': n,
            'Actual': valor_meta(n, valor, unidad),
            'Meta': texto_meta(meta, unidad, mayor).replace('meta ', ''),
            'Estado': chip(estado_meta(valor, meta, mayor)),
        } for n, valor, _, unidad, meta, mayor in METAS]))

    col1, col2 = st.columns([3, 2], gap='medium')
    with col1:
        seccion('Venta particular mensual', f'{ACT} contra el mismo mes de {ANT} (con IVA)')
        filas = []
        for m in MESES_RANGO:
            filas.append({'Mes': MESES[m - 1], 'Período': ACT,
                          'Venta': v_act.loc[v_act['fecha'].dt.month == m, 'total'].sum()})
            filas.append({'Mes': MESES[m - 1], 'Período': ANT,
                          'Venta': v_ant.loc[v_ant['fecha'].dt.month == m, 'total'].sum()})
        actual_vs_anterior(pd.DataFrame(filas), 'Mes', 'Venta', ACT, ANT, None, alto=340)

    with col2:
        # Hallazgos que se escriben solos según los datos del período
        puntos = []
        if crec_ing is not None:
            puntos.append(f'Los ingresos netos de la venta particular {"crecieron" if crec_ing >= 0 else "cayeron"} '
                          f'<b>{pct(abs(crec_ing))}</b> frente al mismo tramo de {ANT} '
                          f'({mm(TOT["ingresos"])} vs {mm(TOT_ANT["ingresos"])}).')
        var_ot, var_tk = variacion(len(s_act), len(s_ant)), variacion(ticket_ot(s_act), ticket_ot(s_ant))
        if var_ot is not None and var_tk is not None:
            puntos.append(f'Las OT particulares {"suben" if var_ot >= 0 else "bajan"} {pct(abs(var_ot))} y el ticket '
                          f'promedio {"sube" if var_tk >= 0 else "baja"} {pct(abs(var_tk))}.')
        er_l = ER_ACT.drop('Total')
        if er_l['ingresos'].sum() > 0:
            top = er_l['margen_directo'].idxmax()
            puntos.append(f'<b>{top}</b> genera el {pct(100 * er_l.loc[top, "margen_directo"] / TOT["margen_directo"])} '
                          f'del margen directo, con un margen de {pct(er_l.loc[top, "pct_md"])}.')
        if not s_act.empty:
            tk_tipo = s_act[s_act['total'] > 0].groupby('origen_tipo')['total'].mean()
            if 'Oftalmólogo/a' in tk_tipo and 'Sin registro' in tk_tipo:
                puntos.append(f'Las OT derivadas por oftalmólogo tienen un ticket de <b>{clp(tk_tipo["Oftalmólogo/a"])}</b>, '
                              f'contra {clp(tk_tipo["Sin registro"])} de las OT sin derivador registrado. '
                              f'Solo el {pct(der_act)} de las OT registra quién deriva.')
        if not pd.isna(dsc_act):
            puntos.append(f'El <b>{pct(dsc_act)}</b> de las OT particulares tiene descuento.')
        if not pd.isna(lt_act):
            puntos.append(f'Una OT particular demora <b>{miles(lt_act)} días</b> (mediana) en entregarse y el '
                          f'<b>{pct(atr_act)}</b> se entrega después de la fecha prometida.')
        lista = ''.join(f'<li>{p}</li>' for p in puntos)
        st.markdown(f'<div class="hallazgos" style="margin-top:22px"><h4>Hallazgos clave del período</h4>'
                    f'<ul>{lista}</ul></div>', unsafe_allow_html=True)


# =============================================================================
# 10. Pestaña: Resultados por línea
# =============================================================================
with tabs[1]:
    ini_k = costos['fecha'].min() if not costos.empty else None
    fin_k = costos['fecha'].max() if not costos.empty else None
    rango_k = f'{ini_k:%d-%m-%Y} al {fin_k:%d-%m-%Y}' if ini_k is not None else 'sin informe cargado'
    nota('<b>Cómo se calcula.</b> Ingresos = boletas de OT particulares, sin IVA. Costo directo = % de costo de '
         f'cada línea según el informe de costos ({rango_k}), aplicado a todo el período: es un <b>estimado</b>. '
         'El costo de los cristales viene de las listas de precios de los proveedores (preliminar, por validar). '
         + ('Gastos operacionales = montos de la barra lateral, repartidos según los ingresos de cada línea.'
            if HAY_GASTOS else 'Gastos operacionales: <b>pendientes</b>.'))

    fila_kpi([
        tarjeta('Ingresos netos', mm(TOT['ingresos']), crec_ing, COMPARADO,
                meta=semaforo('Crecimiento de ingresos')),
        tarjeta('Margen directo', mm(TOT['margen_directo']),
                variacion(TOT['margen_directo'], TOT_ANT['margen_directo']), COMPARADO,
                meta=semaforo('Margen directo')),
        tarjeta('Gastos operacionales', mm(TOT['gastos']) if HAY_GASTOS else 'Pendiente', None,
                f'{clp(GASTO_MENSUAL)} al mes' if HAY_GASTOS else 'se ingresan en la barra lateral'),
        tarjeta('Margen operacional', mm(TOT['margen_operacional']) if HAY_GASTOS else 'Pendiente',
                variacion(TOT['margen_operacional'], TOT_ANT['margen_operacional']) if HAY_GASTOS else None,
                COMPARADO if HAY_GASTOS else 'faltan los gastos de COALIVI',
                meta=semaforo('Margen operacional') if HAY_GASTOS else None),
    ])

    seccion('Estado de resultados por línea de producto', f'período {ACT} · montos netos de IVA')
    gasto_txt = (lambda x: clp(x)) if HAY_GASTOS else (lambda x: 'pendiente')
    tabla_html(pd.DataFrame({
        'Línea': ER_ACT.index,
        'Ingresos netos': ER_ACT['ingresos'].apply(clp),
        '% del total': ER_ACT['pct_ingresos'].apply(pct),
        'Costo directo': ER_ACT['costo'].apply(clp),
        'Margen directo': ER_ACT['margen_directo'].apply(clp),
        '% MD': ER_ACT['pct_md'].apply(pct),
        'Gastos operacionales': ER_ACT['gastos'].apply(gasto_txt),
        'Margen operacional': ER_ACT['margen_operacional'].apply(gasto_txt),
        '% MO': ER_ACT['pct_mo'].apply(pct),
    }), con_total=True, negativas=('Margen operacional', '% MO'))

    col1, col2 = st.columns(2, gap='medium')
    with col1:
        seccion('Margen por línea', '% sobre los ingresos netos')
        mg = ER_ACT.drop('Total')
        mg = mg[mg['ingresos'] > 0].rename_axis('linea').reset_index()
        columnas_mg = ['pct_md', 'pct_mo'] if HAY_GASTOS else ['pct_md']
        mg = mg.melt(id_vars='linea', value_vars=columnas_mg, var_name='Margen', value_name='pct')
        mg['Margen'] = mg['Margen'].map({'pct_md': 'Margen directo', 'pct_mo': 'Margen operacional'})
        fig = px.bar(mg, x='linea', y='pct', color='Margen', barmode='group',
                     color_discrete_map={'Margen directo': AZUL, 'Margen operacional': NARANJO},
                     category_orders={'linea': pr.LINEAS})
        fig.update_traces(marker_line_width=0, hovertemplate='%{x}<br>%{y:.1f}%<extra>%{fullData.name}</extra>')
        fig.add_hline(y=0, line_color='#C3C8D4', line_width=1)
        fig.update_layout(xaxis_title=None, yaxis_title='%')
        mostrar(fig, 360)
    with col2:
        seccion('Ingresos netos por línea', f'{ACT} vs {ANT}')
        il = pd.concat([ER_ACT.drop('Total')[['ingresos']].assign(Período=ACT),
                        ER_ANT.drop('Total')[['ingresos']].assign(Período=ANT)]).rename_axis('Línea').reset_index()
        actual_vs_anterior(il, 'Línea', 'ingresos', ACT, ANT, None, alto=360)

    seccion('Ingresos netos mensuales por línea', f'período {ACT}')
    im = v_act.groupby(['mes', 'linea'], as_index=False)['total'].sum()
    im['ingresos'] = im['total'] / FACTOR_IVA
    im['Mes'] = im['mes'].dt.month.apply(lambda m: MESES[m - 1])
    fig = px.bar(im, x='Mes', y='ingresos', color='linea', color_discrete_map=COLOR_LINEA,
                 category_orders={'linea': pr.LINEAS, 'Mes': MESES})
    fig.update_traces(marker_line_color='white', marker_line_width=1.5,
                      hovertemplate='%{x}<br>$%{y:,.0f}<extra>%{fullData.name}</extra>')
    fig.update_layout(barmode='stack', xaxis_title=None, yaxis_title=None)
    mostrar(fig, 340)

    with st.expander('Supuestos del cálculo'):
        tabla_html(pd.DataFrame({
            'Línea': PCT_COSTO.index,
            '% costo sobre precio de venta (con IVA)': (100 * PCT_COSTO['pct_costo']).apply(pct),
            'Fuente': PCT_COSTO['fuente'],
        }))
        st.caption((f'Gastos operacionales: {clp(GASTO_MENSUAL)} al mes × {miles(MESES_PERIODO, 1)} meses del '
                    f'período = {clp(GASTO_MENSUAL * MESES_PERIODO)}. ' if HAY_GASTOS else
                    'Gastos operacionales: pendientes. ') + f'Precios con IVA: {"sí" if con_iva else "no"}.')

    # ---- Rentabilidad real, OT por OT, del informe de costos
    with st.expander('Rentabilidad real del informe de costos (OT particulares)', expanded=True):
        k_sel = periodo(costos)
        if k_sel.empty:
            k_sel = costos
            st.caption(f'El período elegido no tiene informe de costos: se muestra todo el informe ({rango_k}).')
        else:
            st.caption(f'Productos de OT particulares del {k_sel["fecha"].min():%d-%m-%Y} al '
                       f'{k_sel["fecha"].max():%d-%m-%Y} con venta y costo. Venta neta de IVA; costo neto.')
        ok = k_sel[k_sel['costeable']].copy() if not k_sel.empty else k_sel
        if ok.empty:
            sin_datos('El informe de costos no está cargado.')
        else:
            ok['venta_neta'] = ok['venta'] / FACTOR_IVA

            def margen_por(por, solo_cristal=False):
                base = ok[ok['familia'] == 'Cristal'] if solo_cristal else ok
                r = base.groupby(por).agg(OT=('ot', 'nunique'), venta=('venta_neta', 'sum'), costo=('costo', 'sum'))
                r['margen'] = r['venta'] - r['costo']
                r['pct'] = 100 * r['margen'] / r['venta']
                r['por_ot'] = r['margen'] / r['OT']
                return r.sort_values('venta', ascending=False)

            def tabla_margen(r, nombre):
                tabla_html(pd.DataFrame({nombre: r.index, 'OT': r['OT'].map(miles), 'Venta neta': r['venta'].apply(clp),
                                         'Costo': r['costo'].apply(clp), 'Margen': r['margen'].apply(clp),
                                         '% margen': r['pct'].apply(pct), 'Margen por OT': r['por_ot'].apply(clp)}))

            tot = margen_por(pd.Series('Total', index=ok.index))
            fila_kpi([
                tarjeta('OT con costo', miles(tot['OT'].iloc[0]), None, 'en el informe'),
                tarjeta('Venta neta', mm(tot['venta'].iloc[0]), None, 'productos con costo'),
                tarjeta('Margen directo real', pct(tot['pct'].iloc[0]), None, mm(tot['margen'].iloc[0])),
                tarjeta('Margen por OT', clp(tot['por_ot'].iloc[0]), None, 'neto de IVA'),
            ])
            col1, col2 = st.columns(2, gap='medium')
            with col1:
                seccion('Por familia de producto')
                tabla_margen(margen_por('familia'), 'Familia')
                seccion('Cristales por tipo de lente', 'margen por OT = solo los cristales')
                tabla_margen(margen_por('tipo_lente', solo_cristal=True)
                             .sort_values('por_ot', ascending=False), 'Tipo de lente')
            with col2:
                seccion('Según quién deriva la OT')
                tabla_margen(margen_por('origen_tipo'), 'Derivado por')
                seccion('OT con y sin descuento')
                r = margen_por('con_descuento')
                r.index = r.index.map({True: 'Con descuento', False: 'Sin descuento'})
                tabla_margen(r, 'Descuento')

            prod = ok.groupby('producto').agg(OT=('ot', 'nunique'), venta=('venta_neta', 'sum'),
                                              costo=('costo', 'sum'), familia=('familia', 'first')).reset_index()
            prod = prod[prod['OT'] >= 5]
            prod['margen'] = 100 * (prod['venta'] - prod['costo']) / prod['venta']
            fig = px.scatter(prod, x='venta', y='margen', size='OT', color='familia', hover_name='producto',
                             color_discrete_map=COLOR_FAMILIA, size_max=28,
                             title='Productos: venta neta vs margen (5 o más OT)')
            fig.update_traces(marker=dict(line=dict(color='white', width=1.5), opacity=0.85),
                              hovertemplate='<b>%{hovertext}</b><br>Venta $%{x:,.0f}<br>Margen %{y:.1f}%<extra></extra>')
            fig.update_layout(xaxis_title='venta neta ($)', yaxis_title='margen directo (%)', xaxis_showgrid=True,
                              xaxis_gridcolor=GRILLA)
            mostrar(fig, 380)

            col1, col2 = st.columns(2, gap='medium')
            with col1:
                pv = margen_por('proveedor', solo_cristal=True).reset_index()
                pv['texto'] = pv['pct'].apply(pct) + ' · ' + pv['OT'].map(miles) + ' OT'
                barras_h(pv, 'pct', 'proveedor', 'Margen de cristales por proveedor', formato='%',
                         etiqueta='texto', alto=320)
            with col2:
                def tabla_productos(df):
                    return pd.DataFrame({'Producto': df['producto'], 'OT': df['OT'],
                                         'Venta neta': df['venta'].apply(clp), 'Margen': df['margen'].apply(pct)})
                st.markdown('**Top 10 por venta**')
                st.dataframe(tabla_productos(prod.sort_values('venta', ascending=False).head(10)), hide_index=True)
                st.markdown('**Top 10 de menor margen**')
                st.dataframe(tabla_productos(prod.sort_values('margen').head(10)), hide_index=True)


# =============================================================================
# 11. Pestaña: Comercial (ventas, derivación y descuentos)
# =============================================================================
with tabs[2]:
    fila_kpi([
        tarjeta('Boletas emitidas', miles(boletas_act), variacion(boletas_act, boletas_ant), COMPARADO),
        tarjeta('Ticket promedio por boleta', clp(ticket_act), variacion(ticket_act, ticket_ant), COMPARADO),
        tarjeta('Unidades vendidas', miles(unidades(v_act)), variacion(unidades(v_act), unidades(v_ant)), COMPARADO),
        tarjeta('Venta con IVA', mm(venta_act), variacion(venta_act, venta_ant), COMPARADO),
    ])

    seccion('Tendencia histórica de la venta particular', 'con IVA · la zona sombreada es el período seleccionado')
    hist = ventas.groupby('mes', as_index=False)['total'].sum()
    hist = hist[hist['mes'] <= fecha_max].reset_index(drop=True)
    hist['Mes'] = etiqueta_mes(hist['mes'])
    fig = px.line(hist, x='Mes', y='total', markers=True)
    fig.update_traces(line=dict(color=AZUL, width=2), marker=dict(size=7, color=AZUL),
                      hovertemplate='%{x}<br>$%{y:,.0f}<extra></extra>')
    dentro = hist.index[(hist['mes'] >= ini.replace(day=1)) & (hist['mes'] <= fin)]
    if len(dentro):
        fig.add_vrect(x0=dentro.min() - 0.5, x1=dentro.max() + 0.5, fillcolor=NARANJO, opacity=0.12, line_width=0)
    fig.update_layout(xaxis_title=None, yaxis_title=None, xaxis_type='category')
    mostrar(fig, 300)

    col1, col2 = st.columns(2, gap='medium')
    with col1:
        seccion('OT particulares por mes', f'{ACT} vs {ANT}')
        filas = []
        for m in MESES_RANGO:
            filas.append({'Mes': MESES[m - 1], 'Período': ACT, 'OT': (s_act['fecha'].dt.month == m).sum()})
            filas.append({'Mes': MESES[m - 1], 'Período': ANT, 'OT': (s_ant['fecha'].dt.month == m).sum()})
        actual_vs_anterior(pd.DataFrame(filas), 'Mes', 'OT', ACT, ANT, None, formato=' OT', alto=320)
    with col2:
        seccion('Ticket promedio por OT', f'{ACT} vs {ANT} · con IVA')
        filas = []
        for m in MESES_RANGO:
            filas.append({'Mes': MESES[m - 1], 'Período': ACT, 'Ticket': ticket_ot(s_act[s_act['fecha'].dt.month == m])})
            filas.append({'Mes': MESES[m - 1], 'Período': ANT, 'Ticket': ticket_ot(s_ant[s_ant['fecha'].dt.month == m])})
        actual_vs_anterior(pd.DataFrame(filas), 'Mes', 'Ticket', ACT, ANT, None, alto=320)

    # ---- Derivación desde el Centro Médico (propuesta priorizada de captación)
    with st.expander('Derivación desde el Centro Médico', expanded=True):
        if s_act.empty:
            sin_datos('No hay OT particulares para el período.')
        else:
            tk_tipo_act = s_act[s_act['total'] > 0].groupby('origen_tipo')['total'].mean()
            fila_kpi([
                tarjeta('OT con derivador registrado', pct(der_act), diferencia(der_act, der_ant), COMPARADO,
                        puntos=True, meta=semaforo('OT con derivador registrado')),
                tarjeta('Derivadas por oftalmólogo', pct(100 * (s_act['origen_tipo'] == 'Oftalmólogo/a').mean()),
                        None, 'de las OT particulares'),
                tarjeta('Ticket derivada por oftalmólogo', clp(tk_tipo_act.get('Oftalmólogo/a')), None, 'con IVA'),
                tarjeta('Ticket sin derivador registrado', clp(tk_tipo_act.get('Sin registro')), None, 'con IVA'),
            ])
            col1, col2 = st.columns(2, gap='medium')
            with col1:
                seccion('OT según quién deriva', f'{ACT} vs {ANT}')
                td = pd.concat([s_act.assign(Período=ACT), s_ant.assign(Período=ANT)])
                td = td.groupby(['origen_tipo', 'Período'], as_index=False).size().rename(columns={'size': 'OT'})
                actual_vs_anterior(td, 'origen_tipo', 'OT', ACT, ANT, None, formato=' OT', alto=320,
                                   orden_x=ORDEN_DERIVADOR)
            with col2:
                seccion('Ticket promedio según quién deriva', f'{ACT} vs {ANT} · con IVA')
                tt = pd.concat([s_act.assign(Período=ACT), s_ant.assign(Período=ANT)])
                tt = tt[tt['total'] > 0].groupby(['origen_tipo', 'Período'], as_index=False)['total'].mean()
                actual_vs_anterior(tt, 'origen_tipo', 'total', ACT, ANT, None, alto=320, orden_x=ORDEN_DERIVADOR)
            col1, col2 = st.columns(2, gap='medium')
            with col1:
                og = (s_act[s_act['origen'] != 'Sin registro'].groupby(['origen', 'origen_tipo'])
                      .agg(ot=('ot', 'size'), venta=('total', 'sum')).reset_index()
                      .sort_values('ot', ascending=False).head(12))
                og['texto'] = og['ot'].map(miles) + ' OT · ' + og['origen_tipo']
                barras_h(og, 'ot', 'origen', f'Profesionales que más derivan · {ACT}', formato=' OT',
                         etiqueta='texto', alto=400)
            with col2:
                seccion('% de OT con derivador registrado', 'mes a mes, últimos 24 meses')
                dm = ot_sala[ot_sala['fecha'] > fecha_max - pd.DateOffset(months=24)]
                dm = dm.groupby('mes', as_index=False)['origen_tipo'].apply(lambda x: 100 * (x != 'Sin registro').mean())
                dm.columns = ['mes', 'pct']
                linea_mensual(dm, 'pct', formato='%', alto=360, meta=meta_der, color=SERIES[3])

    with st.expander('Ventas: forma de pago, vendedoras, días y precios'):
        col1, col2 = st.columns(2, gap='medium')
        with col1:
            fp = v_act.groupby('forma_pago', as_index=False)['total'].sum()
            fp['texto'] = (100 * fp['total'] / fp['total'].sum()).apply(pct)
            barras_h(fp, 'total', 'forma_pago', 'Venta por forma de pago', etiqueta='texto', alto=320)
        with col2:
            vd = v_act.groupby('vendedora').agg(venta=('total', 'sum'), boletas=('folio', 'nunique')).reset_index()
            vd['ticket'] = vd['venta'] / vd['boletas']
            vd = vd.sort_values('venta')
            fig = px.bar(vd, x='venta', y='vendedora', orientation='h', custom_data=['boletas', 'ticket'],
                         title='Venta y ticket por vendedora')
            fig.update_traces(marker_color=AZUL, hovertemplate='<b>%{y}</b><br>Venta $%{x:,.0f}<br>'
                              'Boletas %{customdata[0]:,.0f}<br>Ticket $%{customdata[1]:,.0f}<extra></extra>')
            fig.update_layout(xaxis_title=None, yaxis_title=None)
            mostrar(fig, 320)
        col1, col2 = st.columns(2, gap='medium')
        with col1:
            ds = v_act.groupby(v_act['fecha'].dt.dayofweek)['total'].sum().reindex(range(7), fill_value=0)
            ds = pd.DataFrame({'día': DIAS, 'venta': ds.values})
            fig = px.bar(ds, x='día', y='venta', title='Venta por día de la semana')
            fig.update_traces(marker_color=AZUL, hovertemplate='%{x}<br>$%{y:,.0f}<extra></extra>')
            fig.update_layout(xaxis_title=None, yaxis_title=None)
            mostrar(fig, 320)
        with col2:
            def precio_medio(v, tipo):
                x = v[v['tipo'] == tipo]
                u = unidades(x)
                return x['total'].sum() / u if u else np.nan
            productos = [('Armazón', 'Armazón'), ('Cristal (unidad)', 'Cristales'), ('Lentes de sol', 'Gafa')]
            st.markdown('**Precio medio por unidad (con IVA)**')
            st.dataframe(pd.DataFrame({
                'Producto': [p for p, _ in productos],
                ACT: [clp(precio_medio(v_act, t)) for _, t in productos],
                ANT: [clp(precio_medio(v_ant, t)) for _, t in productos],
                'Variación': [pct(variacion(precio_medio(v_act, t), precio_medio(v_ant, t))) for _, t in productos],
            }), hide_index=True)
            nc_act = periodo(ventas_todas)
            nc = 100 * nc_act.loc[nc_act['es_nc'], 'total'].sum() / venta_act if venta_act else np.nan
            st.caption(f'Notas de crédito: {pct(nc, 2)} de la venta particular del período.')

    with st.expander('Descuentos'):
        if s_act.empty:
            sin_datos()
        else:
            con, sin = s_act[s_act['con_descuento']], s_act[~s_act['con_descuento']]
            fila_kpi([
                tarjeta('OT con descuento', pct(dsc_act), diferencia(dsc_act, dsc_ant), COMPARADO,
                        invertir=True, puntos=True),
                tarjeta('Descuento promedio', pct(s_act['descuento_pct'].mean()), None, 'cuando hay descuento'),
                tarjeta('Ticket con descuento', clp(ticket_ot(con)), None, 'con IVA'),
                tarjeta('Ticket sin descuento', clp(ticket_ot(sin)), None, 'con IVA'),
            ])
            col1, col2 = st.columns(2, gap='medium')
            with col1:
                seccion('% de OT con descuento', 'mes a mes, últimos 24 meses')
                dm = ot_sala[ot_sala['fecha'] > fecha_max - pd.DateOffset(months=24)]
                dm = dm.groupby('mes', as_index=False)['con_descuento'].mean()
                dm['pct'] = 100 * dm['con_descuento']
                linea_mensual(dm, 'pct', formato='%', alto=320, color=SERIES[1])
            with col2:
                dd = s_act['descuento_pct'].dropna().astype(int).value_counts().sort_index().reset_index()
                dd.columns = ['descuento', 'OT']
                dd['descuento'] = dd['descuento'].astype(str) + '%'
                fig = px.bar(dd, x='descuento', y='OT', title=f'Descuentos aplicados · {ACT}')
                fig.update_traces(marker_color=AZUL, hovertemplate='%{x}<br>%{y} OT<extra></extra>')
                fig.update_layout(xaxis_title='% de descuento', yaxis_title=None, xaxis_type='category')
                mostrar(fig, 320)


# =============================================================================
# 12. Pestaña: Operaciones
# =============================================================================
with tabs[3]:
    sin_fecha = 100 * s_act['fecha_entrega'].isna().mean() if not s_act.empty else np.nan
    fila_kpi([
        tarjeta('Plazo de entrega (mediana)', f'{miles(lt_act)} días', variacion(lt_act, lt_ant), COMPARADO,
                invertir=True, meta=semaforo('Plazo de entrega (mediana)')),
        tarjeta('Entregas fuera de plazo', pct(atr_act), diferencia(atr_act, atr_ant),
                COMPARADO, invertir=True, puntos=True, meta=semaforo('Entregas fuera de plazo')),
        tarjeta('OT entregadas', miles(s_act['fecha_entrega'].notna().sum()), None, f'de {miles(len(s_act))} OT'),
        tarjeta('OT sin fecha de entrega', pct(sin_fecha), None, 'pendientes o sin registro'),
    ])

    col1, col2 = st.columns(2, gap='medium')
    with col1:
        seccion('Días entre la OT y la entrega', f'OT particulares · período {ACT}')
        lt = s_act['lead_time'].dropna() if not s_act.empty else pd.Series(dtype=float)
        if lt.empty:
            sin_datos()
        else:
            fig = px.histogram(lt.to_frame('días'), x='días', nbins=30)
            fig.update_traces(marker_color=AZUL, marker_line_color='white', marker_line_width=1.5,
                              hovertemplate='%{x} días<br>%{y} OT<extra></extra>')
            fig.add_vline(x=lt.median(), line_color=NARANJO, line_width=2,
                          annotation_text=f'Mediana {miles(lt.median())} días', annotation_position='top right')
            fig.add_vline(x=meta_lt, line_color=GRIS, line_width=1,
                          annotation_text=f'Meta {miles(meta_lt)}', annotation_position='bottom right')
            fig.update_layout(xaxis_title='días', yaxis_title='OT', bargap=0.05)
            mostrar(fig, 340)
    with col2:
        seccion('Plazo de entrega mes a mes', 'mediana de días (últimos 24 meses)')
        lm = ot_sala[ot_sala['fecha'] > fecha_max - pd.DateOffset(months=24)]
        lm = lm.groupby('mes', as_index=False)['lead_time'].median()
        linea_mensual(lm, 'lead_time', formato=' días', alto=340, meta=meta_lt)

    col1, col2 = st.columns(2, gap='medium')
    with col1:
        seccion('Entregas fuera de plazo mes a mes', '% de las OT entregadas (últimos 24 meses)')
        am = ot_sala[ot_sala['fecha'] > fecha_max - pd.DateOffset(months=24)]
        am = am.groupby('mes').apply(pct_atraso).reset_index(name='pct')
        linea_mensual(am, 'pct', formato='%', alto=320, meta=meta_atr, color=NARANJO)
    with col2:
        seccion('OT por día de la semana', f'OT particulares · período {ACT}')
        dsem = s_act.groupby(s_act['fecha'].dt.dayofweek).size().reindex(range(7), fill_value=0)
        dsem = pd.DataFrame({'día': DIAS, 'OT': dsem.values})
        fig = px.bar(dsem, x='día', y='OT')
        fig.update_traces(marker_color=AZUL, hovertemplate='%{x}<br>%{y} OT<extra></extra>')
        fig.update_layout(xaxis_title=None, yaxis_title=None)
        mostrar(fig, 320)

    seccion('Desempeño por vendedora', f'OT particulares · período {ACT}')
    if s_act.empty:
        sin_datos()
    else:
        dv = s_act.groupby('vendedora').agg(
            OT=('ot', 'size'), ticket=('total', lambda x: x[x > 0].mean()),
            derivacion=('origen_tipo', lambda x: 100 * (x != 'Sin registro').mean()),
            descuento=('con_descuento', 'mean'), plazo=('lead_time', 'median')).reset_index()
        dv = dv[dv['OT'] >= 5].sort_values('OT', ascending=False)
        tabla_html(pd.DataFrame({
            'Vendedora': dv['vendedora'], 'OT': dv['OT'].map(miles), 'Ticket promedio': dv['ticket'].apply(clp),
            '% con derivador registrado': dv['derivacion'].apply(pct),
            '% con descuento': (100 * dv['descuento']).apply(pct),
            'Plazo de entrega (mediana)': dv['plazo'].apply(lambda x: f'{miles(x)} días'),
        }))


# =============================================================================
# 13. Pestaña: Calidad de datos
# =============================================================================
with tabs[4]:
    sin_or = 100 - der_act if not pd.isna(der_act) else np.nan
    sin_costo = ((costos['venta'] > 0) & (costos['costo'] == 0)).sum() if not costos.empty else 0
    corregidos = (costos['arreglo'] != '').sum() if 'arreglo' in costos.columns else 0
    fila_kpi([
        tarjeta('Venta con OT mal registrada', pct(err_act), diferencia(err_act, err_ant), 'toda la sala · ' + COMPARADO,
                invertir=True, puntos=True, meta=semaforo('Venta con OT mal registrada')),
        tarjeta('OT sin profesional que deriva', pct(sin_or), None, f'período {ACT}'),
        tarjeta('Productos vendidos sin costo', miles(sin_costo), None,
                f'de {miles(len(costos))} en el informe' if not costos.empty else ''),
        tarjeta('Precios corregidos', miles(corregidos), None, 'separador de miles o cantidad invertida'),
    ])

    col1, col2 = st.columns(2, gap='medium')
    with col1:
        seccion('% del monto con OT mal o no registrada', 'mes a mes · toda la sala (si la OT está mal, no se sabe si es particular)')
        em = ventas_sala.assign(monto_error=np.where(ventas_sala['estado_ot'] != 'Registrada', ventas_sala['total'], 0))
        em = em.groupby('mes', as_index=False)[['monto_error', 'total']].sum()
        em['pct'] = 100 * em['monto_error'] / em['total']
        linea_mensual(em, 'pct', formato='%', alto=320, meta=meta_err)
    with col2:
        seccion('Registro de OT por vendedora', f'% de líneas de boleta con OT mal o no registrada · {ACT} · toda la sala')
        ev = vs_act.assign(error=vs_act['estado_ot'] != 'Registrada').groupby('vendedora').agg(
            lineas=('error', 'size'), errores=('error', 'sum')).reset_index()
        sin_vend = ev[ev['vendedora'].str.lower() == 'sin registro']
        ev = ev[(ev['lineas'] >= 20) & (ev['vendedora'].str.lower() != 'sin registro')]
        ev['pct'] = 100 * ev['errores'] / ev['lineas']
        ev['texto'] = ev['pct'].apply(pct) + ' (' + ev['errores'].astype(int).astype(str) + ' de ' + \
            ev['lineas'].astype(int).map(miles) + ')'
        if ev.empty:
            sin_datos()
        else:
            barras_h(ev, 'pct', 'vendedora', None, formato='%', etiqueta='texto', alto=320)
        if not sin_vend.empty:
            n, e = int(sin_vend['lineas'].iloc[0]), int(sin_vend['errores'].iloc[0])
            st.caption(f'Además, {miles(n)} líneas de boleta no tienen vendedora asignada, y {miles(e)} de '
                       f'ellas tienen la OT mal o no registrada.')

    seccion('Auditoría de calidad', 'hallazgos para pedir correcciones a TI')
    mal_miles = (costos['arreglo'] == 'Separador de miles').sum() if 'arreglo' in costos.columns else 0
    invertidos = (costos['arreglo'] == 'Cantidad y precio invertidos').sum() if 'arreglo' in costos.columns else 0
    ot_sin_entrega = 100 * ot_sala['fecha_entrega'].isna().mean() if not ot_sala.empty else np.nan
    v_cristal = costos[(costos['familia'] == 'Cristal') & (costos['venta'] > 0)] if not costos.empty else costos
    v5 = 100 * (v_cristal['fuente_costo'] == 'Lista proveedores V5 (preliminar)').mean() if not v_cristal.empty else np.nan
    hallazgos = [
        ('Alta', 'OT sala', 'OT particulares sin profesional que deriva (toda la historia)',
         pct(100 - pct_derivacion(ot_sala)) if not ot_sala.empty else '–'),
        ('Alta', 'Boletas', 'Monto con OT mal o no registrada (toda la sala, toda la historia)', pct(pct_error(ventas_sala))),
        ('Alta', 'Costos', 'Precios con separador de miles mal leído (14.99 en vez de 14.990)', f'{miles(mal_miles)} productos'),
        ('Alta', 'Costos', 'Gafas con cantidad y precio invertidos', f'{miles(invertidos)} productos'),
        ('Media', 'Costos', 'Cristales con costo preliminar (lista de proveedores por validar)', pct(v5)),
        ('Media', 'Costos', 'Productos vendidos sin costo (quedan fuera del margen)', f'{miles(sin_costo)} productos'),
        ('Media', 'OT sala', 'OT particulares sin fecha real de entrega', pct(ot_sin_entrega)),
    ]
    color_criticidad = {'Alta': 'No cumple', 'Media': 'En riesgo'}
    tabla_html(pd.DataFrame([{'Criticidad': chip(color_criticidad[c], c), 'Fuente': f, 'Hallazgo': h,
                              'Magnitud': f'<b>{v}</b>'} for c, f, h, v in hallazgos]))

    with st.expander('Trazabilidad de la OT particular entre sistemas'):
        ot_sal = set(ot_sala['ot'].dropna()) if not ot_sala.empty else set()
        ot_bol = set(ventas_todas['ot'].dropna())
        if not costos.empty:
            en_rango = ot_sala[ot_sala['fecha'].between(costos['fecha'].min(), costos['fecha'].max())]
            ot_rango = set(en_rango[en_rango['total'] > 0]['ot'].dropna())
        else:
            ot_rango = set()
        ot_cos = set(costos['ot'].dropna()) if not costos.empty else set()
        cruces = [
            ['OT particulares sin boleta asociada', len(ot_sal - ot_bol), len(ot_sal)],
            ['OT particulares con venta, dentro del período del informe, sin costeo', len(ot_rango - ot_cos), len(ot_rango)],
            ['Líneas de boleta sin OT que se pueda cruzar (no se sabe si son particulares)',
             int((~ventas_sala['ot'].isin(set(datos['ot_sala']['ot'].dropna()))).sum()), len(ventas_sala)],
        ]
        tabla_html(pd.DataFrame([{'Cruce': c, 'Sin match': miles(a), 'Universo': miles(b),
                                  '%': pct(100 * a / b if b else np.nan)} for c, a, b in cruces]))

    with st.expander('Bases cargadas, actualización y descarga de tablas'):
        resumen = []
        for clave, info in pr.BASES.items():
            df = datos[clave]
            tiene_fecha = 'fecha' in df.columns and not df.empty
            resumen.append({
                'Base': info['nombre'], 'Origen': origen_datos[clave], 'Filas': miles(len(df)),
                'Desde': f'{df["fecha"].min():%d-%m-%Y}' if tiene_fecha else '',
                'Hasta': f'{df["fecha"].max():%d-%m-%Y}' if tiene_fecha else '',
            })
        tabla_html(pd.DataFrame(resumen))
        st.markdown("""
**Actualizar solo para esta sesión:** barra lateral → *Cargar / actualizar bases de datos* → subir las
planillas originales (`CONSOLIDADO VENTAS ÓPTICA Y OT + BOLETAS`, `OT APPS FINAL - Sala de ventas particular`,
`informe_costos_Sala_ventas`).

**Actualizar para todos:** correr `python generar_datos.py "carpeta con las planillas"` y subir a GitHub la
carpeta `data` nueva (ver README).
""")
        descargas = {'ventas': ventas_todas, 'ot_sala': ot_sala, 'costos': costos}
        columnas = st.columns(3)
        for n_, (clave, info) in enumerate(pr.BASES.items()):
            df = descargas[clave]
            if not df.empty:
                columnas[n_ % 3].download_button(f'{info["nombre"]} (particular)',
                                                 df.to_csv(index=False).encode('utf-8-sig'),
                                                 file_name=f'coalivi_particular_{clave}.csv', mime='text/csv',
                                                 key=f'descarga_{clave}')
